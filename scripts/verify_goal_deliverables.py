#!/usr/bin/env python3
"""Requirement-level acceptance checks for the AutoRecon agentic goal."""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

from PIL import Image, ImageStat
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
DEMO = ROOT / "demo_runs"
PRESENTATION = ROOT / "presentation"
STEPS = {"extraction", "reconstruction", "registration", "tissue", "roi", "surface"}
CASES = {
    "anhuijing": "2026-07-20T170535Z0000_6f4bbefc",
    "caoli": "2026-07-20T170547Z0000_ff7bd5da",
    "baihu": "2026-07-20T172019Z0000_e47e3b68",
}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)
    print(f"PASS  {message}")


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def check_case(name: str, run_id: str) -> tuple[dict, dict]:
    root = DEMO / run_id
    manifest = load(root / "demo_manifest.json")
    evidence = load(root / "evidence_bank.json")
    require(manifest["status"] == "completed", f"{name}: demo completed")
    for path in manifest["artifacts"].values():
        require(Path(path).is_file(), f"{name}: artifact exists: {Path(path).name}")
    qc = {item["step_id"]: item for item in evidence["step_qc"]}
    require(set(qc) == STEPS, f"{name}: Evidence Bank covers six pipeline steps")
    require(all(len(item["metrics"]) >= 2 for item in qc.values()), f"{name}: every step has >=2 QC metrics")
    require(all(metric.get("source") for item in qc.values() for metric in item["metrics"]), f"{name}: every QC metric has a source")
    require(bool(evidence.get("claims")), f"{name}: Evidence Bank contains claims")
    require(bool(evidence.get("report_limits")), f"{name}: Evidence Bank contains report limits")
    report = (root / "report.md").read_text(encoding="utf-8")
    require("## Per-step QC Evidence" in report, f"{name}: report exposes per-step QC")
    require("## Fetal Brain MRI Professional Report" in report, f"{name}: professional research report exists")
    require("## Capability Boundary" in report, f"{name}: report declares capability boundary")
    return manifest, evidence


def main() -> int:
    tests = subprocess.run(
        ["python", "-m", "unittest", "discover", "-s", "tests"],
        cwd=ROOT, capture_output=True, text=True,
    )
    require(tests.returncode == 0, "automated agentic test suite passes")

    an, an_ev = check_case("AnHuiJing", CASES["anhuijing"])
    cao, cao_ev = check_case("CaoLi", CASES["caoli"])
    bai, bai_ev = check_case("BaiHuPing", CASES["baihu"])
    an_status = {x["step_id"]: x["status"] for x in an_ev["step_qc"]}
    cao_status = {x["step_id"]: x["status"] for x in cao_ev["step_qc"]}
    bai_status = {x["step_id"]: x["status"] for x in bai_ev["step_qc"]}
    require(set(an_status.values()) == {"pass"}, "AnHuiJing demonstrates all-six-step PASS")
    require(cao_status["registration"] == "review", "CaoLi demonstrates registration REVIEW")
    require(bai["plan"]["display_steps"] == ["reconstruction"], "BaiHuPing display scope is reconstruction only")
    require(bai["plan"]["report_steps"] == [], "BaiHuPing report scope consumes all evidence")
    require(bai_status["tissue"] == bai_status["roi"] == "fail", "BaiHuPing exposes Tissue/ROI geometry failures")

    shots = [
        "01_anhuijing_user_flow.png", "02_caoli_qc_review.png",
        "03_baihu_scope_evidence.png", "04_anhuijing_extraction_viewer.png",
        "05_caoli_registration_mpr.png",
    ]
    for name in shots:
        path = PRESENTATION / "screenshots" / name
        require(path.is_file(), f"screenshot exists: {name}")
        with Image.open(path) as image:
            require(image.width >= 1000 and image.height >= 700, f"screenshot has presentation resolution: {name}")
            stat = ImageStat.Stat(image.convert("L"))
            require(stat.stddev[0] > 10, f"screenshot is not blank: {name}")

    ppt = PRESENTATION / "AutoRecon_Agentic_Group_Meeting.zh-CN.pptx"
    pdf = PRESENTATION / "AutoRecon_Agentic_Group_Meeting.zh-CN.pdf"
    require(ppt.is_file() and ppt.stat().st_size > 500_000, "group-meeting PPT exists and contains media")
    deck = Presentation(ppt)
    require(len(deck.slides) == 14, "group-meeting PPT contains 14 slides")
    deck_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text_frame"))
    for phrase in ("按正确顺序运行", "按依赖顺序排列任务", "结构化需求单", "报告里的每句话都能找到依据", "通过只说明技术检查合格"):
        require(phrase in deck_text.replace("“", "").replace("”", ""), f"PPT includes plain-language explanation: {phrase}")
    for phrase in ("三种用户需求", "用户原话", "用户输入", "处理正确性", "只展示 reconstruction", "系统正确分开两种范围"):
        require(phrase in deck_text, f"PPT demonstrates request diversity and correctness: {phrase}")
    require("白名单 DAG" not in deck_text and "Local validator" not in deck_text, "PPT removes unexplained engineering jargon")
    require(pdf.is_file() and pdf.stat().st_size > 500_000, "PPT rendered successfully to final PDF")
    for name in ("PROJECT_COMPLETION_AUDIT.zh-CN.md", "DEMO_CASES.zh-CN.md"):
        require((PRESENTATION / name).is_file(), f"delivery document exists: {name}")

    print("\nALL GOAL DELIVERABLE CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
