#!/usr/bin/env python3
"""Build the audited AutoRecon agentic group-meeting deck."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "AutoRecon_Agentic_Group_Meeting.zh-CN.pptx"
SHOTS = ROOT / "presentation" / "screenshots"

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(246, 248, 246)
PANEL = RGBColor(255, 255, 255)
INK = RGBColor(27, 42, 38)
MUTED = RGBColor(91, 108, 102)
LINE = RGBColor(215, 225, 220)
GREEN = RGBColor(20, 123, 107)
GREEN2 = RGBColor(226, 243, 238)
AMBER = RGBColor(191, 126, 31)
AMBER2 = RGBColor(255, 244, 219)
RED = RGBColor(178, 68, 68)
BLUE = RGBColor(62, 109, 163)
FONT = "Noto Sans CJK SC"


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        x, y, w, h,
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def pill(slide, text, x, y, w, fill=GREEN2, color=GREEN):
    rect(slide, x, y, w, Inches(.34), fill, fill)
    add_text(slide, text, x, y + Inches(.015), w, Inches(.29), 10, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def title(slide, heading, kicker, number):
    add_text(slide, kicker.upper(), Inches(.55), Inches(.31), Inches(8.5), Inches(.25), 9, GREEN, True)
    add_text(slide, heading, Inches(.55), Inches(.58), Inches(11.8), Inches(.55), 25, INK, True)
    add_text(slide, f"{number:02d}", Inches(12.25), Inches(.38), Inches(.5), Inches(.3), 10, MUTED, True, PP_ALIGN.RIGHT)


def footer(slide):
    add_text(slide, "AutoRecon · Agentic fetal MRI research workflow · 2026-07-21",
             Inches(.55), Inches(7.16), Inches(8), Inches(.2), 8, MUTED)


def bullet_list(slide, items, x, y, w, h, size=15, color=INK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.05)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.font.name = FONT; p.font.size = Pt(size); p.font.color.rgb = color
        p.space_after = Pt(8); p.text = "•  " + p.text
    return box


def add_picture_contain(slide, path, x, y, w, h, border=True):
    path = Path(path)
    if border: rect(slide, x, y, w, h, PANEL, LINE)
    with Image.open(path) as im:
        ratio = im.width / im.height
    target = w / h
    if ratio > target:
        pw, ph = w, w / ratio
        px, py = x, y + (h - ph) / 2
    else:
        ph, pw = h, h * ratio
        px, py = x + (w - pw) / 2, y
    return slide.shapes.add_picture(str(path), px, py, pw, ph)


def flow_box(slide, label, sub, x, y, w, accent=GREEN):
    rect(slide, x, y, w, Inches(.86), PANEL, LINE)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Inches(.08), Inches(.86))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    add_text(slide, label, x + Inches(.18), y + Inches(.12), w - Inches(.25), Inches(.25), 13, INK, True)
    add_text(slide, sub, x + Inches(.18), y + Inches(.42), w - Inches(.25), Inches(.28), 9, MUTED)


def arrow(slide, x1, y1, x2, y2, color=GREEN):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color; c.line.width = Pt(2)
    c.line.end_arrowhead = True


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = BG
    footer(slide)
    return slide


def build() -> Path:
    prs = Presentation(); prs.slide_width = W; prs.slide_height = H

    # 1
    s = new_slide(prs)
    add_text(s, "AUTORECON / AGENTIC MRI", Inches(.62), Inches(.58), Inches(5), Inches(.3), 10, GREEN, True)
    add_text(s, "从自然语言需求到\n可追溯研究报告", Inches(.62), Inches(1.25), Inches(6.5), Inches(1.45), 34, INK, True)
    add_text(s, "大模型理解需求 · 六步质量检查 · 证据库 · 用户视角可视化",
             Inches(.65), Inches(2.9), Inches(6.4), Inches(.62), 15, MUTED)
    pill(s, "工程闭环已验收", Inches(.65), Inches(3.73), Inches(1.55))
    pill(s, "研究决策支持", Inches(2.34), Inches(3.73), Inches(1.48), AMBER2, AMBER)
    add_picture_contain(s, SHOTS / "05_caoli_registration_mpr.png", Inches(7.25), Inches(.65), Inches(5.45), Inches(5.95))
    add_text(s, "组会报告 · 2026-07-21", Inches(.65), Inches(6.44), Inches(3.5), Inches(.3), 11, MUTED)

    # 2
    s = new_slide(prs); title(s, "目标、边界与验收结论", "Executive audit", 2)
    cards = [
        ("目标", "用户说出需求 → 系统整理任务 → 按正确顺序运行 → 给出结果和报告", GREEN),
        ("完成", "理解需求、确认计划、运行工具、逐步质检、保存证据、生成报告均已连通", BLUE),
        ("边界", "专业研究型报告，不在缺少常模与专家签署时做临床诊断", AMBER),
    ]
    for i, (h, b, c) in enumerate(cards):
        x = Inches(.65 + i * 4.17); rect(s, x, Inches(1.45), Inches(3.82), Inches(1.55), PANEL, LINE)
        pill(s, h, x + Inches(.18), Inches(1.63), Inches(.78), GREEN2 if c == GREEN else AMBER2 if c == AMBER else RGBColor(231,239,248), c)
        add_text(s, b, x + Inches(.18), Inches(2.08), Inches(3.45), Inches(.7), 14, INK, True)
    add_text(s, "完成度不是由单个绿色测试证明，而由需求级证据共同证明", Inches(.65), Inches(3.45), Inches(11.8), Inches(.38), 19, INK, True)
    bullet_list(s, [
        "自动化测试覆盖：三种范围互不干扰、拒绝未知工具、接口备用方案、密钥保护、六步质检。",
        "3 个真实既有 run 的演示：全通过、配准需复核、下游空间错误被发现。",
        "每份报告都保存原始需求、执行计划、每步质检来源、结论依据和能力限制。",
        "外部 LLM 不接收本次受试者衍生证据；样例采用本地确定性隐私模式。",
    ], Inches(.72), Inches(4.03), Inches(11.8), Inches(2.15), 15)

    # 3
    s = new_slide(prs); title(s, "用户流程：一句需求如何变成可以检查的结果", "User journey", 3)
    labels = [
        ("01 用户需求", "自由中文 / 英文"), ("02 需求整理", "拆出参数和三种范围"),
        ("03 执行计划", "按依赖顺序排列任务"), ("04 开始运行", "用户确认后后台执行"),
        ("05 质量检查", "每步指标与来源"), ("06 最终报告", "只根据已有证据总结"),
    ]
    for i, (a,b) in enumerate(labels):
        x = Inches(.55 + i * 2.11); flow_box(s, a, b, x, Inches(1.65), Inches(1.78), GREEN if i < 3 else BLUE)
        if i < 5: arrow(s, x + Inches(1.8), Inches(2.08), x + Inches(2.08), Inches(2.08))
    rect(s, Inches(.65), Inches(3.08), Inches(12.0), Inches(2.9), PANEL, LINE)
    add_text(s, "关键行为", Inches(.9), Inches(3.35), Inches(1.5), Inches(.3), 13, GREEN, True)
    bullet_list(s, [
        "点击“生成计划”后会显示处理中；计划完成后仍需用户点击“运行”。",
        "运行哪些步骤、页面展示哪些步骤、报告引用哪些步骤，三者互不干扰。",
        "模型只能选择系统已有工具；工具不存在、顺序矛盾、路径虚构或参数缺失时不会运行。",
        "页面持续更新状态，并保留当前图像窗口；每一步显示短日志和质量指标。",
    ], Inches(.9), Inches(3.82), Inches(11.3), Inches(1.85), 15)

    # 4
    s = new_slide(prs); title(s, "处理流程：六个步骤如何前后衔接", "Pipeline", 4)
    names = [("1", "脑提取"), ("2", "三维重建"), ("3", "图像配准"), ("4", "组织分割"), ("5", "脑区分区"), ("6", "表面生成")]
    for i,(n,name) in enumerate(names):
        x=Inches(.55+i*2.1); rect(s,x,Inches(1.55),Inches(1.78),Inches(1.0),GREEN2 if i==2 else PANEL,GREEN if i==2 else LINE)
        add_text(s,n,x+Inches(.12),Inches(1.73),Inches(.3),Inches(.3),12,GREEN,True)
        add_text(s,name,x+Inches(.42),Inches(1.72),Inches(1.25),Inches(.36),13,INK,True)
        if i<5: arrow(s,x+Inches(1.8),Inches(2.05),x+Inches(2.08),Inches(2.05))
    rect(s, Inches(.75), Inches(3.0), Inches(11.75), Inches(2.55), PANEL, LINE)
    add_text(s, "第 3 步生成统一的配准空间", Inches(1.0), Inches(3.3), Inches(4.0), Inches(.35), 20, GREEN, True)
    add_text(s, "outputs/registered_input/<subject>.nii.gz", Inches(1.0), Inches(3.82), Inches(5.0), Inches(.3), 14, INK, True)
    arrow(s, Inches(5.5), Inches(4.0), Inches(7.0), Inches(4.0), GREEN)
    add_text(s, "Tissue + ROI", Inches(7.15), Inches(3.55), Inches(2.0), Inches(.3), 17, INK, True)
    add_text(s, "图像尺寸、方向和位置必须一致", Inches(7.15), Inches(3.95), Inches(3.8), Inches(.3), 12, MUTED)
    arrow(s, Inches(10.2), Inches(4.0), Inches(11.0), Inches(4.0), GREEN)
    add_text(s, "Surface", Inches(11.15), Inches(3.55), Inches(1.0), Inches(.3), 17, INK, True)
    add_text(s, "依赖两者", Inches(11.15), Inches(3.95), Inches(1.0), Inches(.3), 12, MUTED)
    add_text(s, "系统从任务顺序、实际输入路径和图像空间检查三方面保证后续步骤使用配准结果。", Inches(1.0), Inches(4.78), Inches(10.8), Inches(.35), 14, MUTED)

    # 5
    s = new_slide(prs); title(s, "大模型框架：先听懂需求，再安排任务", "LLM workflow", 5)
    flow_box(s,"第一阶段 · 听懂需求","自然语言 → 结构化需求单",Inches(.65),Inches(1.5),Inches(3.0),GREEN)
    flow_box(s,"第二阶段 · 安排任务","需求单 + 可用工具 → 执行顺序",Inches(5.15),Inches(1.5),Inches(3.0),BLUE)
    flow_box(s,"本地自动检查","格式 / 顺序 / 路径 / 环境",Inches(9.65),Inches(1.5),Inches(3.0),AMBER)
    arrow(s,Inches(3.68),Inches(1.93),Inches(5.1),Inches(1.93)); arrow(s,Inches(8.18),Inches(1.93),Inches(9.6),Inches(1.93))
    cols = [
        ("结构化需求单保存什么", ["用户原话和目标", "受试者、目录、孕周、图谱", "运行 / 展示 / 报告范围", "限制、缺失信息、无法满足项"]),
        ("任务安排有什么限制", ["只能选择系统已有工具", "写清每一步的前置步骤", "只使用检查过的本地路径", "计划错误时只让模型修正一次"]),
        ("接口不可用怎么办", ["兼容两类常见大模型接口", "系统自动尝试可用接口", "记住成功的接口类型", "仍失败时使用安全备用规则"]),
    ]
    for i,(h,items) in enumerate(cols):
        x=Inches(.65+i*4.2); rect(s,x,Inches(2.85),Inches(3.75),Inches(2.65),PANEL,LINE)
        add_text(s,h,x+Inches(.2),Inches(3.1),Inches(3.3),Inches(.3),16,INK,True)
        bullet_list(s,items,x+Inches(.18),Inches(3.55),Inches(3.35),Inches(1.6),13)
    add_text(s,"大模型不能随意写命令、发明工具，也不能直接判断医学正常或异常。",Inches(.75),Inches(5.85),Inches(11.8),Inches(.38),16,RED,True,PP_ALIGN.CENTER)

    # 6
    s = new_slide(prs); title(s, "三种范围：运行、页面展示和报告引用互不混淆", "Request scopes", 6)
    scopes=[("运行范围","实际计算哪些步骤","决定计算资源和必须补齐的前置步骤",GREEN),
            ("展示范围","页面展示哪些步骤","只改变页面，不会删除证据",BLUE),
            ("报告范围","报告引用哪些步骤","留空表示使用全部已有证据",AMBER)]
    for i,(a,b,c,col) in enumerate(scopes):
        y=Inches(1.45+i*1.35); rect(s,Inches(.7),y,Inches(5.5),Inches(1.02),PANEL,LINE)
        pill(s,a,Inches(.92),y+Inches(.17),Inches(1.72),GREEN2 if col==GREEN else AMBER2 if col==AMBER else RGBColor(231,239,248),col)
        add_text(s,b,Inches(2.85),y+Inches(.12),Inches(3.0),Inches(.3),15,INK,True)
        add_text(s,c,Inches(2.85),y+Inches(.52),Inches(3.0),Inches(.28),11,MUTED)
    rect(s,Inches(6.65),Inches(1.45),Inches(5.95),Inches(3.72),PANEL,LINE)
    add_text(s,"一句话如何被理解",Inches(6.95),Inches(1.75),Inches(2.5),Inches(.3),16,GREEN,True)
    add_text(s,"“只展示 reconstruction，\n但生成完整报告”",Inches(6.95),Inches(2.25),Inches(4.9),Inches(.75),24,INK,True)
    add_text(s,'页面：只显示“重建”\n报告：仍使用全部步骤证据',Inches(6.95),Inches(3.28),Inches(4.8),Inches(.72),17,BLUE,True)
    add_text(s,"因此隐藏页面上的某一步，不会把该步的重要问题从报告中删掉。",Inches(6.95),Inches(4.3),Inches(4.8),Inches(.45),12,MUTED)

    # 7
    s = new_slide(prs); title(s, "证据库：报告里的每句话都能找到依据", "Evidence Bank", 7)
    chain=[("用户需求","原话 + 整理后的需求"),("执行计划","步骤顺序 + 三种范围"),("每步质检","数值 + 标准 + 来源"),("报告结论","结论 + 依据 + 可信程度"),("使用限制","不能回答什么 + 待复核项")]
    for i,(a,b) in enumerate(chain):
        x=Inches(.55+i*2.52); flow_box(s,a,b,x,Inches(1.5),Inches(2.1),GREEN if i<3 else AMBER)
        if i<4: arrow(s,x+Inches(2.12),Inches(1.93),x+Inches(2.48),Inches(1.93))
    rect(s,Inches(.7),Inches(2.95),Inches(12.0),Inches(2.65),PANEL,LINE)
    add_text(s,"保存到本地的记录",Inches(.95),Inches(3.23),Inches(2.5),Inches(.3),15,GREEN,True)
    add_text(s,"intent_spec.json   task_plan.json   evidence_bank.json   medical_evidence.json   report.md",Inches(.95),Inches(3.72),Inches(11.0),Inches(.4),16,INK,True)
    bullet_list(s,["每个指标都保存数值、单位、判断标准、结果和来源文件。","每条结论都保存证据和可信程度；报告同时说明不能过度解释的部分。","大模型的分析会进入报告正文，用户不需要翻隐藏日志。"],Inches(.95),Inches(4.3),Inches(10.9),Inches(1.05),13)

    # 8
    s = new_slide(prs); title(s, "六步质量检查：每一步检查什么、结果怎么看", "Quality checks", 8)
    rows=[("01 脑提取","3","掩膜是否生成 / 范围 / 切片覆盖","通过"),
          ("02 三维重建","4","图像可读 / 有效体素 / 强度范围 / 大小","通过"),
          ("03 图像配准","3","Mattes / GC / 配准图像是否存在","通过 / 需复核"),
          ("04 组织分割","3","标签 / 前景范围 / 是否处于配准空间","通过 / 未通过"),
          ("05 脑区分区","3","标签 / 前景范围 / 是否处于配准空间","通过 / 未通过"),
          ("06 表面生成","2","必要表面文件 / 统计文件是否齐全","通过")]
    headers=["处理步骤","指标数","系统自动检查的内容","案例中出现过的结果"]
    xs=[.65,3.25,4.4,10.25]; ws=[2.5,1.05,5.7,2.35]
    for x,w,hdr in zip(xs,ws,headers): add_text(s,hdr,Inches(x),Inches(1.4),Inches(w),Inches(.28),9,MUTED,True)
    for i,row in enumerate(rows):
        y=1.8+i*.72; rect(s,Inches(.58),Inches(y),Inches(12.05),Inches(.58),PANEL,LINE)
        for x,w,val in zip(xs,ws,row): add_text(s,val,Inches(x),Inches(y+.13),Inches(w),Inches(.25),12,INK,True if x in (.65,10.25) else False)
    add_text(s,"“通过”只说明技术检查合格，不代表医学正常；配准结果仍需人工对照图像复核。",Inches(.68),Inches(6.3),Inches(11.8),Inches(.3),12,AMBER,True)

    # 9
    s = new_slide(prs); title(s, "三种用户需求：系统不靠一个固定模板回答", "Request diversity", 9)
    cases = [
        ("A · 全流程查看", "“展示完整六步、每步质量指标，\n并生成完整中文报告。”", "导入已有结果 · 展示全部 · 报告全部", "六步全部通过", GREEN),
        ("B · 发现需复核项", "“展示全部步骤、质量指标和日志，\n生成完整中文研究报告。”", "导入已有结果 · 保留自动复核结论", "配准需要人工复核", AMBER),
        ("C · 页面只看一步", "“只展示 reconstruction，\n但用全部证据生成完整报告。”", "页面只看重建 · 报告仍检查六步", "发现组织和脑区空间问题", BLUE),
    ]
    for i,(name,request,understanding,result,col) in enumerate(cases):
        x=Inches(.55+i*4.18); rect(s,x,Inches(1.35),Inches(3.86),Inches(5.35),PANEL,LINE)
        pill(s,name,x+Inches(.25),Inches(1.65),Inches(1.55),GREEN2 if col==GREEN else AMBER2 if col==AMBER else RGBColor(231,239,248),col)
        add_text(s,"用户原话",x+Inches(.25),Inches(2.25),Inches(1.2),Inches(.25),11,MUTED,True)
        add_text(s,request,x+Inches(.25),Inches(2.62),Inches(3.35),Inches(1.05),16,INK,True)
        add_text(s,"系统理解",x+Inches(.25),Inches(4.0),Inches(1.1),Inches(.25),11,col,True)
        add_text(s,understanding,x+Inches(.25),Inches(4.35),Inches(3.25),Inches(.58),13,INK,True)
        add_text(s,"结果证据",x+Inches(.25),Inches(5.2),Inches(1.1),Inches(.25),11,col,True)
        add_text(s,result,x+Inches(.25),Inches(5.57),Inches(3.25),Inches(.58),15,col,True)
    add_text(s,"多样性来自用户目标和范围不同；正确性由计划、质量指标、来源文件和最终报告共同证明。",Inches(.75),Inches(6.85),Inches(11.8),Inches(.3),13,GREEN,True,PP_ALIGN.CENTER)

    # 10
    s = new_slide(prs); title(s, "案例 A · AnHuiJing：完整查看和完整报告", "Complete review", 10)
    rect(s,Inches(.6),Inches(1.25),Inches(12.05),Inches(.82),GREEN2,GREEN2)
    add_text(s,"用户输入",Inches(.85),Inches(1.48),Inches(1.0),Inches(.25),11,GREEN,True)
    add_text(s,"“导入 AnHuiJing 现有结果，展示完整六步可视化和每一步质量指标，基于证据库生成完整中文专业研究报告。”",Inches(1.75),Inches(1.4),Inches(10.45),Inches(.42),15,INK,True)
    add_picture_contain(s,SHOTS/"01_anhuijing_user_flow.png",Inches(.6),Inches(2.25),Inches(7.45),Inches(4.45))
    rect(s,Inches(8.3),Inches(2.25),Inches(4.35),Inches(4.45),PANEL,LINE)
    pill(s,"处理正确性",Inches(8.6),Inches(2.55),Inches(1.25))
    add_text(s,"系统没有省略步骤",Inches(8.6),Inches(3.05),Inches(3.3),Inches(.35),19,INK,True)
    bullet_list(s,["理解为：导入已有结果","页面：展示全部六步","报告：使用全部六步证据","实际结果：六步质量检查均通过","配准 Mattes −0.33959，GC −0.546999"],Inches(8.55),Inches(3.55),Inches(3.4),Inches(2.2),13)
    add_text(s,"需求、计划、证据库和报告范围一致",Inches(8.6),Inches(6.05),Inches(3.3),Inches(.35),13,GREEN,True)

    # 11
    s = new_slide(prs); title(s, "案例 A · 用户继续逐张查看脑提取结果", "Interactive viewer", 11)
    add_picture_contain(s,SHOTS/"04_anhuijing_extraction_viewer.png",Inches(.65),Inches(1.25),Inches(8.2),Inches(5.65))
    rect(s,Inches(9.1),Inches(1.25),Inches(3.55),Inches(5.65),PANEL,LINE)
    add_text(s,"用户可以做什么",Inches(9.4),Inches(1.62),Inches(2.8),Inches(.3),17,GREEN,True)
    bullet_list(s,["同时查看原始 MRI 和脑掩膜","用滚轮或滑条逐张浏览","点击播放前先加载当前序列","以稳定速度逐张播放","页面刷新状态时不会重置图像窗口"],Inches(9.35),Inches(2.15),Inches(2.8),Inches(2.4),14)
    add_text(s,"质量指标不是孤立分数：用户可以回到对应图像逐张确认。",Inches(9.4),Inches(5.15),Inches(2.75),Inches(.9),15,INK,True)

    # 12
    s = new_slide(prs); title(s, "案例 B · CaoLi：程序跑完，不代表质量一定通过", "Review case", 12)
    rect(s,Inches(.6),Inches(1.25),Inches(12.05),Inches(.82),AMBER2,AMBER2)
    add_text(s,"用户输入",Inches(.85),Inches(1.48),Inches(1.0),Inches(.25),11,AMBER,True)
    add_text(s,"“导入 CaoLi 配准空间结果，展示全部步骤、质量指标和日志，并生成完整中文研究报告。”",Inches(1.75),Inches(1.4),Inches(10.45),Inches(.42),15,INK,True)
    add_picture_contain(s,SHOTS/"05_caoli_registration_mpr.png",Inches(.6),Inches(2.25),Inches(7.45),Inches(4.45))
    rect(s,Inches(8.3),Inches(2.25),Inches(4.35),Inches(4.45),PANEL,LINE)
    pill(s,"处理正确性",Inches(8.6),Inches(2.55),Inches(1.25),AMBER2,AMBER)
    add_text(s,"系统保留了“需要复核”",Inches(8.6),Inches(3.05),Inches(3.3),Inches(.35),19,INK,True)
    bullet_list(s,["理解为：导入并展示全部结果","程序状态：已完成","配准质量：需要人工复核","其余五步：技术检查通过","报告要求查看三方向配准对照图"],Inches(8.55),Inches(3.55),Inches(3.4),Inches(2.2),13)
    add_text(s,"没有把“有结果文件”误写成“配准可靠”",Inches(8.6),Inches(6.05),Inches(3.35),Inches(.35),13,AMBER,True)

    # 13
    s = new_slide(prs); title(s, "案例 C · BaiHuPing：只看一步，报告仍查全部", "Display vs report", 13)
    rect(s,Inches(.6),Inches(1.25),Inches(12.05),Inches(.82),RGBColor(231,239,248),RGBColor(231,239,248))
    add_text(s,"用户输入",Inches(.85),Inches(1.48),Inches(1.0),Inches(.25),11,BLUE,True)
    add_text(s,"“只展示 reconstruction，但用全部可用证据生成完整中文报告，并明确列出需要复核的质量问题。”",Inches(1.75),Inches(1.4),Inches(10.45),Inches(.42),15,INK,True)
    add_picture_contain(s,SHOTS/"03_baihu_scope_evidence.png",Inches(.6),Inches(2.25),Inches(7.45),Inches(4.45))
    rect(s,Inches(8.3),Inches(2.25),Inches(4.35),Inches(4.45),PANEL,LINE)
    pill(s,"处理正确性",Inches(8.6),Inches(2.55),Inches(1.25),RGBColor(231,239,248),BLUE)
    add_text(s,"系统正确分开两种范围",Inches(8.6),Inches(3.05),Inches(3.3),Inches(.35),19,INK,True)
    bullet_list(s,["页面：只显示三维重建","报告：仍检查全部六步","组织分割：空间不一致","脑区分区：空间不一致","最终报告明确列出两个问题"],Inches(8.55),Inches(3.55),Inches(3.4),Inches(2.2),13)
    add_text(s,"用户少看页面，不会导致系统少检查证据",Inches(8.6),Inches(6.05),Inches(3.35),Inches(.35),13,BLUE,True)

    # 14
    s = new_slide(prs); title(s, "系统现在能做什么，下一步还需要什么", "Current scope", 14)
    left=["目前已经做到","理解自然语言并安全安排任务","逐步质量检查、看图和实时日志","证据库与专业研究报告","重要问题不会因页面隐藏而漏报"]
    right=["临床转化仍需","版本化解剖标签映射","孕周常模与百分位曲线","专家视觉 QC 评分规范","独立验证、审计与签署"]
    for i,(items,col) in enumerate(((left,GREEN),(right,AMBER))):
        x=Inches(.7+i*6.15); rect(s,x,Inches(1.4),Inches(5.75),Inches(4.65),PANEL,LINE)
        add_text(s,items[0],x+Inches(.3),Inches(1.75),Inches(4.9),Inches(.35),20,col,True)
        bullet_list(s,items[1:],x+Inches(.3),Inches(2.42),Inches(4.9),Inches(2.5),16)
    rect(s,Inches(1.3),Inches(6.25),Inches(10.75),Inches(.58),GREEN2,GREEN2)
    add_text(s,"结论：系统已可用于受控研究编排与证据审阅；临床诊断能力仍需独立医学验证。",Inches(1.5),Inches(6.38),Inches(10.35),Inches(.28),15,GREEN,True,PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    return OUT


if __name__ == "__main__":
    build()
