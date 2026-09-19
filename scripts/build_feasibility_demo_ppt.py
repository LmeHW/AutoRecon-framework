#!/usr/bin/env python3
"""Build a concise, case-specific deck for feasible/infeasible AutoRecon requests."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from build_group_meeting_ppt import (
    AMBER, AMBER2, BG, BLUE, GREEN, GREEN2, INK, LINE, MUTED, PANEL, RED,
    SHOTS, W, H, add_picture_contain, add_text, arrow, bullet_list, flow_box,
    footer, pill, rect, title,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "AutoRecon_AnHuiJing_Feasibility_Demo.zh-CN.pptx"
PALE_BLUE = RGBColor(231, 239, 248)
PALE_RED = RGBColor(251, 235, 235)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    footer(slide)
    return slide


def status_card(slide, x, y, number, label, state, color=GREEN):
    rect(slide, x, y, Inches(1.72), Inches(1.02), PANEL, LINE)
    add_text(slide, number, x + Inches(.12), y + Inches(.12), Inches(.34), Inches(.28), 11, color, True)
    add_text(slide, label, x + Inches(.43), y + Inches(.11), Inches(1.14), Inches(.3), 13, INK, True)
    pill(slide, state, x + Inches(.43), y + Inches(.54), Inches(.83), GREEN2 if color == GREEN else AMBER2, color)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    # 1 — cover
    s = new_slide(prs)
    add_text(s, "AUTORECON · DEMO DESIGN", Inches(.65), Inches(.55), Inches(4), Inches(.3), 10, GREEN, True)
    add_text(s, "什么能跑？\n什么必须拒绝？", Inches(.65), Inches(1.35), Inches(5.8), Inches(1.35), 35, INK, True)
    add_text(s, "31 周 · AnHuiJing · 原始 BTFE → 六步 Pipeline → 可追溯报告",
             Inches(.68), Inches(3.05), Inches(5.9), Inches(.7), 16, MUTED)
    pill(s, "可行性", Inches(.68), Inches(4.05), Inches(1.1))
    pill(s, "拒绝机制", Inches(1.93), Inches(4.05), Inches(1.25), AMBER2, AMBER)
    pill(s, "逐句证据", Inches(3.33), Inches(4.05), Inches(1.25), PALE_BLUE, BLUE)
    add_picture_contain(s, SHOTS / "01_anhuijing_user_flow.png", Inches(6.75), Inches(.65), Inches(5.9), Inches(5.95))

    # 2 — concrete request
    s = new_slide(prs); title(s, "可行 Demo：输入已经足够具体", "AnHuiJing / request", 2)
    rect(s, Inches(.65), Inches(1.33), Inches(4.1), Inches(4.95), PANEL, LINE)
    pill(s, "用户输入", Inches(.93), Inches(1.63), Inches(1.0))
    add_text(s, "31 周 AnHuiJing\n原始 BTFE NIfTI", Inches(.93), Inches(2.15), Inches(3.3), Inches(.75), 23, INK, True)
    bullet_list(s, ["ROI：31w atlas warp", "Surface：完整生成", "六步日志 + QC + 交互图", "全部证据 → 中文报告"],
                Inches(.92), Inches(3.18), Inches(3.35), Inches(2.3), 15)
    add_text(s, "路径 / GA / atlas / 输出目录均已给出", Inches(.93), Inches(5.72), Inches(3.2), Inches(.3), 12, GREEN, True)
    labels = [("01", "Extraction"), ("02", "Reconstruction"), ("03", "Registration"),
              ("04", "Tissue"), ("05", "ROI warp"), ("06", "Surface")]
    for i, (n, label) in enumerate(labels):
        x = Inches(5.15 + (i % 3) * 2.42); y = Inches(1.58 + (i // 3) * 1.52)
        status_card(s, x, y, n, label, "计划中")
        if i % 3 < 2: arrow(s, x + Inches(1.74), y + Inches(.5), x + Inches(2.34), y + Inches(.5))
    rect(s, Inches(5.15), Inches(4.82), Inches(6.55), Inches(1.23), GREEN2, GREEN2)
    add_text(s, "结论", Inches(5.43), Inches(5.08), Inches(.8), Inches(.28), 13, GREEN, True)
    add_text(s, "运行条件齐全；可生成白名单 DAG", Inches(6.22), Inches(5.04), Inches(4.8), Inches(.35), 20, INK, True)

    # 3 — plan audit
    s = new_slide(prs); title(s, "这份 Plan：主链正确，但有一处自相矛盾", "Plan audit", 3)
    add_text(s, "✓", Inches(.72), Inches(1.4), Inches(.5), Inches(.5), 28, GREEN, True)
    add_text(s, "依赖关系正确", Inches(1.22), Inches(1.48), Inches(2.5), Inches(.35), 20, INK, True)
    flow_box(s, "Input", "路径 + GA + atlas", Inches(.75), Inches(2.08), Inches(2.1), GREEN)
    flow_box(s, "3D pipeline", "六步顺序正确", Inches(3.45), Inches(2.08), Inches(2.1), GREEN)
    flow_box(s, "QC / Metrics", "随步骤生成证据", Inches(6.15), Inches(2.08), Inches(2.1), GREEN)
    flow_box(s, "Dashboard", "日志 + 图像 + 指标", Inches(8.85), Inches(2.08), Inches(2.1), GREEN)
    for x in (2.88, 5.58, 8.28): arrow(s, Inches(x), Inches(2.51), Inches(x + .52), Inches(2.51))
    rect(s, Inches(.75), Inches(3.55), Inches(11.7), Inches(1.75), PALE_RED, RED)
    pill(s, "需修正", Inches(1.02), Inches(3.88), Inches(1.05), PALE_RED, RED)
    add_text(s, "unfulfilled_requests 说“DAG 不生成报告”", Inches(2.27), Inches(3.82), Inches(5.4), Inches(.38), 18, INK, True)
    add_text(s, "但执行器会在 Pipeline 后调用 reporter → report.md", Inches(2.27), Inches(4.34), Inches(6.8), Inches(.35), 16, BLUE, True)
    add_text(s, "因此：可运行 ≠ Plan 文本完全正确", Inches(.8), Inches(5.82), Inches(11.5), Inches(.38), 22, AMBER, True, PP_ALIGN.CENTER)

    # 4 — current truth
    s = new_slide(prs); title(s, "Demo 展示必须区分：计划、运行中、已完成", "Live state", 4)
    add_picture_contain(s, SHOTS / "04_anhuijing_extraction_viewer.png", Inches(.62), Inches(1.3), Inches(7.35), Inches(5.45))
    rect(s, Inches(8.25), Inches(1.3), Inches(4.4), Inches(5.45), PANEL, LINE)
    pill(s, "当前事实", Inches(8.56), Inches(1.63), Inches(1.15), AMBER2, AMBER)
    add_text(s, "agentic_state = running", Inches(8.55), Inches(2.17), Inches(3.4), Inches(.4), 20, INK, True)
    bullet_list(s, ["Extraction：已有日志 / QC / 图", "Reconstruction：已有日志 / 图", "Registration → Surface：尚未完成", "最终报告：尚不能宣称六步结论"],
                Inches(8.5), Inches(2.88), Inches(3.55), Inches(2.25), 14)
    add_text(s, "组会措辞", Inches(8.55), Inches(5.45), Inches(1.1), Inches(.3), 12, BLUE, True)
    add_text(s, "“完整链路已启动，非六步已完成”", Inches(8.55), Inches(5.85), Inches(3.35), Inches(.55), 17, BLUE, True)

    # 5 — infeasible outputs
    s = new_slide(prs); title(s, "不可行需求：系统应停止，并明确说明缺什么", "Fail closed", 5)
    examples = [
        ("缺 atlas", "“做 31w ROI warp”\n但 atlas 路径为空", "missing", "atlas", RED),
        ("缺输入", "subject_dir 不存在\n或无 BTFE_*.nii.gz", "missing", "subject_dir / BTFE", RED),
        ("依赖矛盾", "只跑 surface\n不提供 Tissue / ROI", "warnings", "surface prerequisites", AMBER),
    ]
    for i, (head, req, field, value, color) in enumerate(examples):
        x = Inches(.62 + i * 4.2)
        rect(s, x, Inches(1.35), Inches(3.82), Inches(4.85), PANEL, LINE)
        pill(s, head, x + Inches(.25), Inches(1.66), Inches(1.0), PALE_RED if color == RED else AMBER2, color)
        add_text(s, req, x + Inches(.25), Inches(2.2), Inches(3.2), Inches(.72), 17, INK, True)
        add_text(s, "系统输出", x + Inches(.25), Inches(3.18), Inches(1.0), Inches(.25), 11, MUTED, True)
        rect(s, x + Inches(.25), Inches(3.55), Inches(3.28), Inches(1.25), BG, LINE)
        add_text(s, f'"{field}": [\n  "{value}"\n]', x + Inches(.42), Inches(3.73), Inches(2.9), Inches(.85), 14, color, True)
        pill(s, "不启动 GPU", x + Inches(.25), Inches(5.34), Inches(1.3), PALE_RED, RED)
    add_text(s, "原则：无法验证输入或依赖 → status = failed / plan blocked", Inches(.7), Inches(6.47), Inches(11.8), Inches(.3), 15, RED, True, PP_ALIGN.CENTER)

    # 6 — unsupported asks
    s = new_slide(prs); title(s, "另一类不可行：数据能跑，但问题不能回答", "Unsupported interpretation", 6)
    rect(s, Inches(.65), Inches(1.35), Inches(5.25), Inches(4.95), PANEL, LINE)
    pill(s, "用户问", Inches(.95), Inches(1.68), Inches(.95), PALE_RED, RED)
    add_text(s, "“请判断胎儿脑是否正常，\n并给出临床诊断。”", Inches(.95), Inches(2.25), Inches(4.15), Inches(.9), 25, INK, True)
    add_text(s, "即使六步全部 PASS，也不能推出医学正常", Inches(.95), Inches(3.58), Inches(4.25), Inches(.55), 17, RED, True)
    bullet_list(s, ["无孕周常模曲线", "无异常检测器", "无专家签署"], Inches(.95), Inches(4.55), Inches(3.6), Inches(1.2), 14)
    rect(s, Inches(6.25), Inches(1.35), Inches(6.4), Inches(4.95), GREEN2, GREEN)
    pill(s, "应输出", Inches(6.58), Inches(1.68), Inches(1.0))
    add_text(s, '"unsupported_reason":\n"不能进行独立临床诊断或正常性判断"',
             Inches(6.58), Inches(2.28), Inches(5.25), Inches(1.05), 19, INK, True)
    add_text(s, "仍可提供", Inches(6.58), Inches(3.75), Inches(1.2), Inches(.3), 12, GREEN, True)
    bullet_list(s, ["技术 QC 结果", "可视化与定量结果", "明确的人工复核清单"], Inches(6.58), Inches(4.22), Inches(4.7), Inches(1.35), 16)

    # 7 — knowledge base anatomy
    s = new_slide(prs); title(s, "“知识库”具体是什么：两类证据，不是一个黑盒", "Evidence architecture", 7)
    rect(s, Inches(.65), Inches(1.38), Inches(5.75), Inches(4.95), PANEL, LINE)
    pill(s, "病例证据库", Inches(.95), Inches(1.7), Inches(1.35))
    add_text(s, "evidence_bank.json", Inches(.95), Inches(2.22), Inches(3.8), Inches(.35), 21, INK, True)
    bullet_list(s, ["原始需求与 Plan", "每步 status", "metric：值 / 单位 / 阈值", "source：具体 NIfTI / CSV / surf", "limitations：能说明什么"],
                Inches(.95), Inches(2.9), Inches(4.45), Inches(2.45), 15)
    rect(s, Inches(6.75), Inches(1.38), Inches(5.9), Inches(4.95), PANEL, LINE)
    pill(s, "医学边界库", Inches(7.05), Inches(1.7), Inches(1.35), PALE_BLUE, BLUE)
    add_text(s, "medical_evidence.json", Inches(7.05), Inches(2.22), Inches(4.0), Inches(.35), 21, INK, True)
    bullet_list(s, ["claim_id + statement", "evidence + confidence", "允许怎样改写", "禁止诊断 / 编造标签", "缺少常模与专家签署"],
                Inches(7.05), Inches(2.9), Inches(4.55), Inches(2.45), 15)
    add_text(s, "前者回答“本病例发生了什么” · 后者约束“这些事实能被解释到哪一步”",
             Inches(.7), Inches(6.52), Inches(11.8), Inches(.3), 15, BLUE, True, PP_ALIGN.CENTER)

    # 8 — sentence trace 1
    s = new_slide(prs); title(s, "报告逐句溯源：一句结论，对应一条证据链", "Sentence → source", 8)
    rows = [
        ("报告句 ①", "本病例包含 6 个 BTFE stacks，GA=31 周。", "case_context", "staging 文件计数 + 请求元数据"),
        ("报告句 ②", "Extraction 技术检查通过。", "step_qc.extraction", "success=1.0；coverage=0.282292"),
        ("报告句 ③", "Registration 指标达到自动阈值。", "step_qc.registration", "Mattes=-0.33959；GC=-0.546999"),
        ("报告句 ④", "Tissue 与 registered MRI 几何一致。", "step_qc.tissue", "registered_geometry_match=true"),
    ]
    for i, (tag, sentence, key, evidence) in enumerate(rows):
        y = Inches(1.28 + i * 1.28)
        rect(s, Inches(.62), y, Inches(12.05), Inches(1.02), PANEL, LINE)
        pill(s, tag, Inches(.84), y + Inches(.16), Inches(1.0), PALE_BLUE, BLUE)
        add_text(s, sentence, Inches(2.0), y + Inches(.14), Inches(5.1), Inches(.3), 16, INK, True)
        add_text(s, key, Inches(7.32), y + Inches(.13), Inches(2.05), Inches(.28), 11, GREEN, True)
        add_text(s, evidence, Inches(7.32), y + Inches(.52), Inches(4.75), Inches(.26), 12, MUTED)
    add_text(s, "注：这些是已完成 AnHuiJing_run02 的真实证据；不能冒充当前 running run 的结果。",
             Inches(.7), Inches(6.55), Inches(11.8), Inches(.3), 13, AMBER, True, PP_ALIGN.CENTER)

    # 9 — sentence trace 2 / conclusion
    s = new_slide(prs); title(s, "安全句也要有依据：报告不把 PASS 写成“正常”", "Boundary trace", 9)
    rect(s, Inches(.68), Inches(1.38), Inches(12.0), Inches(1.42), GREEN2, GREEN)
    add_text(s, "报告句 ⑤", Inches(.98), Inches(1.72), Inches(1.2), Inches(.3), 12, GREEN, True)
    add_text(s, "“该流程支持科研定量和专家复核，不支持独立诊断或正常性判断。”",
             Inches(2.1), Inches(1.63), Inches(9.6), Inches(.48), 21, INK, True)
    add_text(s, "依据：interpretation_boundary · 缺 GA 常模 / atlas 标签映射 / 异常检测器 / 专家签署",
             Inches(2.1), Inches(2.25), Inches(9.2), Inches(.28), 12, MUTED)
    cards = [("可运行", "输入与依赖齐全", GREEN), ("可复核", "每句连到 source", BLUE), ("会拒绝", "缺输入 / 越界解释", RED)]
    for i, (head, body, color) in enumerate(cards):
        x = Inches(.72 + i * 4.17)
        rect(s, x, Inches(3.25), Inches(3.82), Inches(2.05), PANEL, LINE)
        pill(s, head, x + Inches(.27), Inches(3.57), Inches(1.05), GREEN2 if color == GREEN else PALE_BLUE if color == BLUE else PALE_RED, color)
        add_text(s, body, x + Inches(.27), Inches(4.22), Inches(3.15), Inches(.4), 19, INK, True)
    add_text(s, "Demo 成功标准：不是“所有卡片变绿”，而是“系统只说证据允许说的话”。",
             Inches(.72), Inches(5.88), Inches(11.8), Inches(.5), 21, INK, True, PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    return OUT


if __name__ == "__main__":
    build()
